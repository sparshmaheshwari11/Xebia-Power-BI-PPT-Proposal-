import os, json, re, uuid
from pathlib import Path
from typing import Optional
import requests
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, UploadFile, HTTPException
from fastapi.responses import HTMLResponse, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pypdf import PdfReader
from docx import Document
from .pptx_builder import build_presentation, validate_proposal
from .prompt import SYSTEM_PROMPT, build_user_prompt

BASE = Path(__file__).resolve().parent.parent
load_dotenv(BASE / ".env")
OUT = BASE / "outputs"
OUT.mkdir(exist_ok=True)
app = FastAPI(title="Xebia Power BI Proposal Studio")
app.mount("/static", StaticFiles(directory=str(BASE / "frontend")), name="static")

def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".txt":
            return data.decode("utf-8", errors="replace")
        if ext == ".pdf":
            from io import BytesIO
            reader = PdfReader(BytesIO(data))
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        if ext == ".docx":
            from io import BytesIO
            doc = Document(BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(paragraphs)
    except Exception as exc:
        raise HTTPException(400, f"Could not read uploaded file: {exc}")
    raise HTTPException(400, "Upload a .txt, .pdf or .docx file.")

def groq_generate(client_name: str, industry: str, requirements: str, tone: str) -> dict:
    api_key = os.getenv("GROQ_API_KEY", "").strip()
    model = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b").strip()
    if not api_key or api_key.startswith("replace_"):
        raise HTTPException(503, "Groq is not configured. Add GROQ_API_KEY to your .env file and restart the server.")
    user_prompt = build_user_prompt(client_name, industry, requirements, tone)
    try:
        response = requests.post(
            "https://api.groq.com/openai/v1/chat/completions",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json={
                "model": model,
                "temperature": 0.35,
                "max_tokens": 16000,
                "response_format": {"type": "json_object"},
                "messages": [
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_prompt}
                ]
            },
            timeout=180
        )
    except requests.RequestException as exc:
        raise HTTPException(502, f"Could not reach Groq API: {exc}")
    if response.status_code >= 400:
        detail = response.text[:1200]
        raise HTTPException(502, f"Groq API returned HTTP {response.status_code}: {detail}")
    try:
        content = response.json()["choices"][0]["message"]["content"]
        data = json.loads(content)
    except Exception as exc:
        raise HTTPException(502, f"Groq returned a response that could not be parsed as proposal JSON: {exc}")
    validate_proposal(data)
    return data

@app.get("/", response_class=HTMLResponse)
def home():
    return (BASE / "frontend" / "index.html").read_text(encoding="utf-8")

@app.get("/api/health")
def health():
    return {
        "status": "ok",
        "provider": "Groq",
        "model": os.getenv("GROQ_MODEL", "openai/gpt-oss-120b"),
        "configured": bool(os.getenv("GROQ_API_KEY", "").strip() and not os.getenv("GROQ_API_KEY", "").startswith("replace_"))
    }

@app.post("/api/generate")
async def generate(
    client_name: str = Form("Client"),
    industry: str = Form("Not specified"),
    tone: str = Form("Executive consulting"),
    requirements_text: str = Form(""),
    file: Optional[UploadFile] = File(None)
):
    source_text = (requirements_text or "").strip()
    has_file = bool(file and file.filename)
    if has_file and source_text:
        raise HTTPException(400, "Choose one input method: upload a BRD or enter plain text, not both.")
    source = source_text
    if has_file:
        data = await file.read()
        limit_mb = int(os.getenv("MAX_UPLOAD_MB", "25"))
        if len(data) > limit_mb * 1024 * 1024:
            raise HTTPException(413, f"File exceeds {limit_mb} MB.")
        source += "\n\n--- UPLOADED REQUIREMENTS ---\n" + extract_text(file.filename, data)
    if len(source.strip()) < 80:
        raise HTTPException(400, "Choose one input method and provide at least 80 characters: upload a BRD, or enter requirements in plain text.")
    if len(source) > 100000:
        source = source[:100000]
    proposal = groq_generate((client_name or "Client").strip() or "Client", (industry or "Not specified").strip() or "Not specified", source, (tone or "Executive consulting").strip() or "Executive consulting")
    output_path = OUT / f"proposal_{uuid.uuid4().hex[:10]}.pptx"
    build_presentation(proposal, output_path)
    return {"download_url": f"/api/download/{output_path.name}", "filename": output_path.name,
            "slide_count": len(proposal["client_slides"]) + 8,
            "title": proposal.get("proposal_title", "Power BI Solution Proposal"),
            "quality_summary": proposal.get("quality_summary", [])}

@app.get("/api/download/{filename}")
def download(filename: str):
    safe = Path(filename).name
    path = OUT / safe
    if not path.exists() or path.suffix.lower() != ".pptx":
        raise HTTPException(404, "Proposal not found.")
    return FileResponse(path, filename=safe, media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
